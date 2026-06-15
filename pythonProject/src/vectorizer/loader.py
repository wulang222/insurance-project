"""文件发现与内容加载 — 按扩展名过滤，提取纯文本。"""

from __future__ import annotations

import os
import fnmatch
import logging
from pathlib import Path
from typing import Any

from vectorizer.config import VectorizerConfig

logger = logging.getLogger(__name__)


# ── 文件扫描 ──

def scan_files(config: VectorizerConfig) -> list[Path]:
    """递归扫描 root_dir，返回所有匹配扩展名且不被排除的文件路径列表。"""
    root = Path(config.scan_root_dir).resolve()
    if not root.is_dir():
        logger.warning("扫描目录不存在: %s", root)
        return []

    matched: list[Path] = []
    for dirpath, dirnames, filenames in os.walk(root):
        cur = Path(dirpath)

        # 排除目录
        dirnames[:] = [d for d in dirnames if d not in config.exclude_dirs]

        for fn in filenames:
            fpath = cur / fn
            ext = fpath.suffix.lower()

            # 扩展名过滤
            if ext not in config.include_extensions:
                continue

            # 文件名通配排除
            excluded = False
            for pat in config.exclude_files:
                if fnmatch.fnmatch(fn, pat):
                    excluded = True
                    break
            if excluded:
                continue

            matched.append(fpath)

    # 限制文件数量
    if config.max_files > 0:
        matched = matched[: config.max_files]

    logger.info("扫描完成，发现 %d 个文件", len(matched))
    return matched


# ── 文件内容加载 ──

def load_file(file_path: Path) -> dict[str, Any] | None:
    """加载单个文件，返回 {"path", "file_type", "content", "metadata"}。
    无法解析的文件返回 None。
    """
    ext = file_path.suffix.lower()
    try:
        loader = _LOADER_MAP.get(ext, _load_text)
        text, meta = loader(file_path)
        return {
            "path": str(file_path),
            "file_type": ext.lstrip("."),
            "content": text,
            "metadata": meta,
        }
    except Exception as e:
        logger.warning("加载失败 [%s]: %s", file_path, e)
        return None


def _load_text(file_path: Path) -> tuple[str, dict]:
    """纯文本类文件：.md .txt .py .java .xml .yaml .json .html .css .js …"""
    content = file_path.read_text(encoding="utf-8", errors="replace")
    return content, {"lines": content.count("\n") + 1, "size_bytes": len(content.encode("utf-8"))}


def _load_pdf(file_path: Path) -> tuple[str, dict]:
    """PDF 文件，逐页提取文字。"""
    try:
        import pypdf
    except ImportError:
        logger.warning("pypdf 未安装，跳过 PDF: %s", file_path)
        return "", {}

    reader = pypdf.PdfReader(str(file_path))
    pages = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages.append(text)
    content = "\n".join(pages)
    return content, {"page_count": len(reader.pages), "size_bytes": file_path.stat().st_size}


def _load_docx(file_path: Path) -> tuple[str, dict]:
    """Word 文档。"""
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx 未安装，跳过 DOCX: %s", file_path)
        return "", {}

    doc = Document(str(file_path))
    paras = [p.text for p in doc.paragraphs]
    content = "\n".join(paras)
    return content, {"paragraphs": len(paras), "size_bytes": file_path.stat().st_size}


def _load_xlsx(file_path: Path) -> tuple[str, dict]:
    """Excel 工作簿，每个单元格内容拼成表格文本。"""
    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl 未安装，跳过 XLSX: %s", file_path)
        return "", {}

    wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    parts: list[str] = []
    sheet_names: list[str] = []
    for sheet_name in wb.sheetnames:
        sheet_names.append(sheet_name)
        ws = wb[sheet_name]
        rows_text: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            rows_text.append(" | ".join(cells))
        parts.append(f"=== Sheet: {sheet_name} ===\n" + "\n".join(rows_text))
    content = "\n\n".join(parts)
    return content, {"sheets": sheet_names, "size_bytes": file_path.stat().st_size}


# 扩展名 → 加载函数映射
_LOADER_MAP: dict[str, callable] = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".doc": _load_docx,
    ".xlsx": _load_xlsx,
    ".xls": _load_xlsx,
}
